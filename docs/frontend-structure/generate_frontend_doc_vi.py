from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "docs" / "frontend-structure"
DOCX_PATH = OUT_DIR / "Frontend_Structure_TuteClass.docx"
IMG_FLOW = OUT_DIR / "frontend-runtime-flow.png"
IMG_TREE = OUT_DIR / "frontend-folder-tree-vi.png"
IMG_LAYER = OUT_DIR / "frontend-layer-usage.png"
IMG_API = OUT_DIR / "frontend-api-data-flow.png"

FONT_DIR = Path("C:/Windows/Fonts")
FONT_REGULAR = FONT_DIR / "arial.ttf"
FONT_BOLD = FONT_DIR / "arialbd.ttf"


def image_font(size, bold=False):
    try:
        return ImageFont.truetype(str(FONT_BOLD if bold else FONT_REGULAR), size=size)
    except OSError:
        return ImageFont.load_default()


TITLE = image_font(39, True)
H1 = image_font(26, True)
H2 = image_font(21, True)
BODY = image_font(18)
SMALL = image_font(15)
CODE = image_font(16)


def wrap_text(draw, text, max_width, font):
    words = text.split()
    lines = []
    current = ""
    for word in words:
        candidate = word if not current else f"{current} {word}"
        if draw.textbbox((0, 0), candidate, font=font)[2] <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def rounded(draw, xy, fill, outline="#D0D5DD", radius=16, width=2):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def draw_box(draw, xy, title, description, fill="#FFFFFF", title_color="#0B2F66"):
    x1, y1, x2, y2 = xy
    rounded(draw, xy, fill, radius=18)
    draw.text((x1 + 20, y1 + 16), title, font=H1, fill=title_color)
    y = y1 + 58
    for raw_line in description.split("\n"):
        for line in wrap_text(draw, raw_line, x2 - x1 - 40, BODY):
            draw.text((x1 + 20, y), line, font=BODY, fill="#475467")
            y += 26


def centered_text(draw, text, xy, font, fill="#0B2F66", spacing=5):
    x1, y1, x2, y2 = xy
    lines = []
    for raw_line in text.split("\n"):
        lines.extend(wrap_text(draw, raw_line, x2 - x1 - 28, font) or [""])
    heights = [draw.textbbox((0, 0), line, font=font)[3] for line in lines]
    total = sum(heights) + spacing * (len(lines) - 1)
    y = y1 + ((y2 - y1) - total) / 2
    for line, height in zip(lines, heights):
        bbox = draw.textbbox((0, 0), line, font=font)
        width = bbox[2] - bbox[0]
        draw.text((x1 + ((x2 - x1) - width) / 2, y), line, font=font, fill=fill)
        y += height + spacing


def arrow(draw, start, end, color="#667085"):
    draw.line([start, end], fill=color, width=3)
    sx, sy = start
    ex, ey = end
    if abs(ex - sx) >= abs(ey - sy):
        points = [(ex, ey), (ex - 12, ey - 7), (ex - 12, ey + 7)] if ex >= sx else [(ex, ey), (ex + 12, ey - 7), (ex + 12, ey + 7)]
    else:
        points = [(ex, ey), (ex - 7, ey - 12), (ex + 7, ey - 12)] if ey >= sy else [(ex, ey), (ex - 7, ey + 12), (ex + 7, ey + 12)]
    draw.polygon(points, fill=color)


def create_runtime_flow_image():
    image = Image.new("RGB", (1600, 920), "#F7FAFC")
    draw = ImageDraw.Draw(image)
    draw.text((55, 34), "Luồng khởi động và hiển thị màn hình", font=TITLE, fill="#0B2F66")
    draw.text((55, 84), "Từ entry point TypeScript đến page, feature và thành phần dùng chung", font=BODY, fill="#475467")

    top_boxes = [
        ((45, 170, 260, 310), "main.tsx", "Gắn React vào DOM\nNạp Ant Design và CSS"),
        ((305, 170, 520, 310), "App.tsx", "Component gốc\nGhép provider và router"),
        ((565, 170, 850, 310), "AppProviders", "BrowserRouter\nQueryClientProvider"),
        ((895, 170, 1120, 310), "router.tsx", "Đọc URL\nLazy-load màn hình"),
        ((1165, 170, 1555, 310), "layouts", "Khung Public, Student,\nTeacher và Admin"),
    ]
    for xy, title, description in top_boxes:
        fill = {"main.tsx": "#FFF7E8", "App.tsx": "#E8F0FF", "AppProviders": "#E8F7F0", "router.tsx": "#F3ECFF", "layouts": "#FFF0E2"}[title]
        draw_box(draw, xy, title, description, fill)

    for first, second in zip(top_boxes, top_boxes[1:]):
        a = first[0]
        b = second[0]
        arrow(draw, (a[2], (a[1] + a[3]) // 2), (b[0], (b[1] + b[3]) // 2))

    draw_box(draw, (150, 500, 535, 695), "pages", "Màn hình gắn với URL\nVí dụ: TeacherCalendarPage.tsx", "#FFFFFF")
    draw_box(draw, (610, 500, 995, 695), "features", "Component và logic nghiệp vụ\nVí dụ: teaching, tuition", "#FFFFFF")
    draw_box(draw, (1070, 500, 1455, 695), "shared", "API client, DTO, hook, UI,\nstyle và tiện ích dùng chung", "#FFFFFF")

    arrow(draw, (1360, 310), (340, 500), "#FF5C00")
    arrow(draw, (535, 598), (610, 598), "#FF5C00")
    arrow(draw, (995, 598), (1070, 598), "#FF5C00")

    rounded(draw, (190, 785, 1410, 860), "#0B2F66", "#0B2F66")
    centered_text(draw, "URL chọn layout và page; page lắp ghép feature; feature dùng nền tảng từ shared.", (220, 798, 1380, 848), H2, "#FFFFFF")
    image.save(IMG_FLOW)


def create_folder_tree_image():
    lines = [
        "frontend/",
        "  package.json            thư viện và câu lệnh npm",
        "  tsconfig.json           cấu hình TypeScript strict",
        "  vite.config.ts          cấu hình dev/build và chia bundle",
        "  openapi/openapi.json    contract API mẫu",
        "  src/",
        "    main.tsx              entry point",
        "    vite-env.d.ts         kiểu cho biến môi trường Vite",
        "    app/",
        "      App.tsx | router.tsx",
        "      layouts/            Public | Student | Teacher | Admin",
        "      providers/          App | Auth | Theme",
        "    pages/",
        "      public/             landing và xác thực",
        "      student/            overview, schedule, assignments...",
        "      teacher/            overview, calendar, tuition...",
        "      admin/              dashboard quản trị",
        "    features/",
        "      teaching/           ScheduleCalendar.tsx",
        "      learning-progress/  LearningProgressChart.tsx",
        "      tuition/            TuitionStatisticsChart.tsx",
        "      assignments/ classes/ auth/ communication/...",
        "    shared/",
        "      api/generated/      schema.d.ts sinh từ OpenAPI",
        "      services/           apiClient.ts dùng Axios",
        "      hooks/              hook dùng chung",
        "      components/         UI dùng chung",
        "      styles/ icons/ utils/",
        "    assets/               ảnh, logo, media khi phát sinh",
    ]

    image = Image.new("RGB", (1600, 1160), "#0B1220")
    draw = ImageDraw.Draw(image)
    draw.text((55, 34), "Cây thư mục frontend hiện tại", font=TITLE, fill="#FFFFFF")
    draw.text((55, 84), "Tên file và vai trò chính sau khi chuyển sang React + TypeScript", font=BODY, fill="#B8C0CC")

    y = 140
    for line in lines:
        stripped = line.strip()
        indent = len(line) - len(line.lstrip())
        x = 65 + indent * 17
        if stripped.endswith("/"):
            color, used_font = "#78D3FF", H2
        elif ".ts" in stripped or ".json" in stripped:
            color, used_font = "#F8D66D", CODE
        else:
            color, used_font = "#D8DEE9", CODE
        draw.text((x, y), line, font=used_font, fill=color)
        y += 34
    image.save(IMG_TREE)


def create_layer_usage_image():
    image = Image.new("RGB", (1600, 900), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    draw.text((55, 34), "Pages, features và shared phối hợp như thế nào?", font=TITLE, fill="#0B2F66")
    draw.text((55, 84), "Ví dụ trực tiếp từ màn hình lịch dạy của giáo viên", font=BODY, fill="#475467")

    draw_box(draw, (120, 155, 1480, 305), "Page", "TeacherCalendarPage.tsx chịu trách nhiệm bố cục màn hình và truyền dữ liệu sự kiện.", "#F3F6FA")
    draw_box(draw, (120, 410, 720, 630), "Feature", "features/teaching/components/\nScheduleCalendar.tsx chứa FullCalendar và hành vi lịch.", "#FFF7E8")
    draw_box(draw, (880, 410, 1480, 630), "Shared", "shared cung cấp API client, DTO, hook,\ncomponent UI và style dùng lại toàn hệ thống.", "#E8F7F0")
    arrow(draw, (650, 305), (420, 410), "#FF5C00")
    arrow(draw, (950, 305), (1180, 410), "#FF5C00")

    rounded(draw, (220, 735, 1380, 830), "#0B2F66", "#0B2F66")
    centered_text(draw, "Page quyết định hiển thị gì; feature xử lý nghiệp vụ; shared cung cấp nền tảng tái sử dụng.", (250, 750, 1350, 815), H2, "#FFFFFF")
    image.save(IMG_LAYER)


def create_api_flow_image():
    image = Image.new("RGB", (1600, 930), "#F8FAFC")
    draw = ImageDraw.Draw(image)
    draw.text((55, 34), "Luồng dữ liệu API và DTO", font=TITLE, fill="#0B2F66")
    draw.text((55, 84), "React Query quản lý trạng thái; Axios gửi request; OpenAPI bảo đảm kiểu dữ liệu", font=BODY, fill="#475467")

    boxes = [
        ((45, 190, 275, 335), "Page/Feature", "Gọi custom hook\nhoặc query"),
        ((335, 190, 605, 335), "React Query", "Cache, loading, error,\nrefetch và retry"),
        ((665, 190, 925, 335), "Axios", "apiClient.ts\nbaseURL + timeout"),
        ((985, 190, 1240, 335), ".NET API", "Endpoint và\nSwagger/OpenAPI"),
        ((1300, 190, 1555, 335), "Dữ liệu", "Response có kiểu\ntrả về React"),
    ]
    for xy, title, description in boxes:
        draw_box(draw, xy, title, description, "#FFFFFF")
    for first, second in zip(boxes, boxes[1:]):
        a, b = first[0], second[0]
        arrow(draw, (a[2], (a[1] + a[3]) // 2), (b[0], (b[1] + b[3]) // 2), "#2F7CFF")

    draw_box(draw, (200, 570, 545, 730), "Swagger/OpenAPI", "Contract do backend xuất ra\n/swagger/v1/swagger.json", "#FFF7E8")
    draw_box(draw, (635, 570, 965, 730), "openapi-typescript", "Lệnh npm generate:api\nchuyển schema thành type", "#F3ECFF")
    draw_box(draw, (1055, 570, 1400, 730), "schema.d.ts", "DTO TypeScript được import\nbởi API service", "#E8F7F0")
    arrow(draw, (545, 650), (635, 650), "#FF5C00")
    arrow(draw, (965, 650), (1055, 650), "#FF5C00")
    arrow(draw, (1225, 570), (800, 335), "#FF5C00")

    draw.text((210, 815), "Kết quả: nếu backend đổi contract, bước generate/check sẽ giúp frontend phát hiện sai kiểu sớm.", font=H2, fill="#0B2F66")
    image.save(IMG_API)


def set_run_font(run, size=None, bold=None, color=None, name="Arial"):
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), name)


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    for run in paragraph.runs:
        set_run_font(run, color=(11, 47, 102))
    return paragraph


def add_para(doc, text, bold_prefix=None):
    paragraph = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, 10.5, True)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest, 10.5)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, 10.5)
    paragraph.paragraph_format.space_after = Pt(6)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        set_run_font(run, 10.5)
        paragraph.paragraph_format.space_after = Pt(3)


def add_numbered(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        run = paragraph.add_run(item)
        set_run_font(run, 10.5)
        paragraph.paragraph_format.space_after = Pt(3)


def add_code(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    set_run_font(run, 9, name="Consolas")
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F3F6FA")
    paragraph._p.get_or_add_pPr().append(shading)
    return paragraph


def add_caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    set_run_font(run, 9, color=(71, 84, 103))
    run.italic = True


def add_picture(doc, path, caption):
    doc.add_picture(str(path), width=Inches(6.7))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_caption(doc, caption)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    set_run_font(run, 9.5, bold=bold)


def add_definition_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["Thành phần", "Định nghĩa và mục đích", "Ví dụ / cách sử dụng"]
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, True)
        shade(table.rows[0].cells[index], "D9EAF7")
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)
    for row_data in rows:
        row = table.add_row()
        row_properties = row._tr.get_or_add_trPr()
        cannot_split = OxmlElement("w:cantSplit")
        row_properties.append(cannot_split)
        cells = row.cells
        for index, value in enumerate(row_data):
            set_cell_text(cells[index], value, bold=(index == 0))
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return table


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("TÀI LIỆU CẤU TRÚC FRONTEND")
    set_run_font(run, 20, True, (11, 47, 102))
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("TuteClass - React + TypeScript + Vite")
    set_run_font(run, 13, color=(71, 84, 103))

    add_para(doc, "Tài liệu này mô tả cấu trúc frontend hiện tại, định nghĩa và mục đích của từng thành phần, luồng chạy của ứng dụng, luồng gọi API và ví dụ đặt code. Thành viên trong nhóm có thể dùng tài liệu như quy ước chung khi phát triển tính năng mới.")

    add_heading(doc, "1. Tổng quan công nghệ")
    add_definition_table(doc, [
        ("React + TypeScript", "React xây dựng giao diện theo component; TypeScript kiểm tra kiểu dữ liệu trước khi chạy.", "Mọi source dùng .ts/.tsx và được kiểm tra bằng npm run typecheck."),
        ("Vite", "Dev server và công cụ build frontend.", "vite.config.ts cấu hình port, React plugin, lazy chunk và vendor chunk."),
        ("Ant Design", "Bộ component UI và hệ thống theme.", "ConfigProvider nằm trong ThemeProvider; Button/Card được dùng ở PagePlaceholder."),
        ("TanStack React Query", "Quản lý server state: cache, loading, error, retry, refetch.", "QueryClientProvider nằm trong AppProviders; useApiHealth dùng useQuery."),
        ("Axios", "HTTP client dùng chung để gọi backend.", "shared/services/apiClient.ts cấu hình baseURL, timeout và kiểu response."),
        ("FullCalendar React", "Hiển thị và tương tác với lịch theo ngày, tuần, tháng.", "ScheduleCalendar.tsx được dùng bởi trang lịch giáo viên và học sinh."),
        ("Recharts", "Hiển thị biểu đồ responsive cho dữ liệu thống kê.", "LearningProgressChart và TuitionStatisticsChart."),
        ("OpenAPI TypeScript", "Sinh DTO TypeScript từ Swagger/OpenAPI của backend.", "npm run generate:api tạo shared/api/generated/schema.d.ts."),
    ])

    add_heading(doc, "2. Luồng chạy của ứng dụng")
    add_picture(doc, IMG_FLOW, "Hình 1. Luồng khởi động và render màn hình")
    add_para(doc, "Luồng bắt đầu tại main.tsx. File này render React vào phần tử #root, nạp CSS reset của Ant Design và CSS toàn cục. App.tsx sau đó ghép AppProviders với AppRouter.")
    add_code(doc, "main.tsx\n  -> App.tsx\n    -> AppProviders\n      -> BrowserRouter\n      -> QueryClientProvider\n      -> ThemeProvider / Ant Design ConfigProvider\n      -> AuthProvider\n    -> router.tsx\n      -> Layout theo vai trò\n        -> Page theo URL\n          -> Feature + Shared")
    add_para(doc, "Ví dụ: URL /teacher/calendar được router lazy-load TeacherCalendarPage.tsx, đặt page trong TeacherLayout và page sử dụng ScheduleCalendar thuộc features/teaching.")

    add_heading(doc, "3. Cây thư mục hiện tại")
    add_picture(doc, IMG_TREE, "Hình 2. Cây thư mục frontend sau khi chuyển sang TypeScript")
    add_para(doc, "Cấu trúc chia code theo trách nhiệm. app quản lý vận hành toàn ứng dụng; pages quản lý màn hình theo URL; features quản lý nghiệp vụ; shared quản lý nền tảng dùng chung.")

    add_heading(doc, "4. Thành phần cấp gốc")
    add_definition_table(doc, [
        ("package.json", "Khai báo dependency và các câu lệnh phát triển, kiểm tra, build, sinh DTO.", "npm run dev, typecheck, build, generate:api và check:api."),
        ("tsconfig.json", "Quy tắc biên dịch TypeScript. strict=true giúp phát hiện null, sai props và sai response type.", "Không tắt strict để né lỗi; sửa đúng type tại nguồn."),
        ("vite.config.ts", "Cấu hình Vite và cách chia bundle.", "Calendar và chart được lazy-load; React và UI library được tách vendor chunk."),
        (".env / vite-env.d.ts", "Cấu hình môi trường và kiểu cho biến Vite.", "VITE_API_BASE_URL xác định địa chỉ backend."),
        ("openapi/openapi.json", "Contract API mẫu khi backend chưa sẵn sàng.", "Dùng npm run generate:api:local; khi backend chạy thì dùng Swagger thật."),
    ])

    add_heading(doc, "5. Thư mục app")
    add_definition_table(doc, [
        ("main.tsx", "Entry point duy nhất để khởi động React.", "Chỉ đặt bootstrap, global CSS và ReactDOM.createRoot; không đặt nghiệp vụ."),
        ("app/App.tsx", "Component gốc kết nối provider và router.", "Giữ file nhỏ để nhìn nhanh cấu trúc cấp cao nhất."),
        ("app/router.tsx", "Bảng ánh xạ URL sang layout và page; hỗ trợ lazy loading.", "Thêm route /teacher/reports tại đây và trỏ tới page tương ứng."),
        ("app/layouts", "Khung giao diện lặp lại theo vai trò.", "TeacherLayout chứa topbar/navigation; các teacher page chỉ cung cấp nội dung Outlet."),
        ("app/providers", "Context và cấu hình toàn ứng dụng.", "AppProviders ghép router/query; ThemeProvider cấu hình Ant Design; AuthProvider cung cấp session."),
    ])

    add_heading(doc, "6. Pages, features và shared")
    add_picture(doc, IMG_LAYER, "Hình 3. Quan hệ giữa page, feature và shared")
    add_definition_table(doc, [
        ("pages", "Mỗi file biểu diễn một màn hình gắn với URL. Page quyết định bố cục và lắp ghép module.", "TeacherOverviewPage dùng LearningProgressChart; không tự chứa code vẽ chart."),
        ("features", "Module theo nghiệp vụ; có thể chứa components, hooks, services và type riêng.", "features/teaching chứa ScheduleCalendar; features/tuition chứa biểu đồ học phí."),
        ("shared", "Hạ tầng và thành phần tái sử dụng ở nhiều nghiệp vụ.", "apiClient, schema DTO, StatusBadge, PagePlaceholder, style token và utility."),
        ("assets", "Tài nguyên tĩnh không phải source code.", "Đặt logo, ảnh minh họa hoặc media vào assets khi phát sinh."),
    ])
    add_para(doc, "Quy tắc phân biệt:", bold_prefix="Quy tắc phân biệt:")
    add_bullets(doc, [
        "Nếu file đại diện cho một URL, đặt trong pages.",
        "Nếu file chỉ phục vụ một nghiệp vụ, đặt trong feature tương ứng.",
        "Nếu file thật sự được dùng bởi nhiều feature, đặt trong shared.",
        "Không chuyển component vào shared chỉ vì có khả năng sẽ dùng lại trong tương lai.",
    ])

    add_heading(doc, "7. Luồng gọi API và sử dụng DTO")
    add_picture(doc, IMG_API, "Hình 4. Luồng React Query, Axios và DTO sinh từ OpenAPI")
    add_para(doc, "API client đọc VITE_API_BASE_URL và tạo Axios instance dùng chung. Custom hook gọi API bằng React Query để nhận cache, trạng thái loading/error và cơ chế hủy request. Kiểu response được lấy từ schema.d.ts sinh bởi OpenAPI.")
    add_code(doc, "// API service có kiểu response sinh từ OpenAPI\ntype SystemInfoDto = components['schemas']['SystemInfoResponse'];\n\napiClient.get<SystemInfoDto>('/api/system/info');\n\n// Hook quản lý server state\nuseQuery({\n  queryKey: ['system-info'],\n  queryFn: ({ signal }) => getSystemInfo(signal),\n});")
    add_para(doc, "Khi backend đã chạy và xuất Swagger tại /swagger/v1/swagger.json, dùng npm run generate:api. Không sửa thủ công schema.d.ts vì file này được sinh tự động.")

    add_heading(doc, "8. Ví dụ các thư viện trong cấu trúc")
    add_definition_table(doc, [
        ("Ant Design", "Chuẩn hóa Button, Card, Form, Table, Modal và theme.", "Component dùng chung có thể import Button/Card; token màu cấu hình tại ThemeProvider."),
        ("FullCalendar", "Feature lịch chịu trách nhiệm plugin tháng/tuần/ngày và sự kiện click.", "TeacherCalendarPage truyền events và editable=true; StudentSchedulePage dùng chế độ chỉ xem."),
        ("React Query + Axios", "React Query quản lý dữ liệu phía client; Axios chỉ thực hiện HTTP request.", "Hook gọi service Axios; page không gọi axios trực tiếp."),
        ("Recharts", "Biểu đồ đặt trong feature sở hữu nghiệp vụ và bọc ResponsiveContainer.", "learning-progress hiển thị điểm/chuyên cần; tuition hiển thị đã thu/còn thiếu."),
    ])

    add_heading(doc, "9. Ví dụ thêm một màn hình mới")
    add_para(doc, "Giả sử cần thêm màn hình báo cáo tiến độ cho giáo viên:")
    add_numbered(doc, [
        "Tạo pages/teacher/TeacherProgressReportPage.tsx để lắp ghép màn hình.",
        "Tạo component nghiệp vụ trong features/learning-progress/components nếu chưa có.",
        "Tạo service trong features/learning-progress/services để gọi endpoint báo cáo.",
        "Tạo hook React Query trong features/learning-progress/hooks để quản lý dữ liệu.",
        "Dùng DTO từ shared/api/generated/schema.d.ts cho request và response.",
        "Khai báo route trong app/router.tsx và để TeacherLayout cung cấp navigation.",
        "Chạy npm run typecheck và npm run build trước khi tạo pull request.",
    ])
    add_code(doc, "pages/teacher/TeacherProgressReportPage.tsx\n  -> features/learning-progress/hooks/useProgressReport.ts\n  -> features/learning-progress/services/progressApi.ts\n  -> shared/services/apiClient.ts\n  -> shared/api/generated/schema.d.ts")

    add_heading(doc, "10. Quy tắc làm việc nhóm")
    add_bullets(doc, [
        "Không dùng lại .js/.jsx; source mới phải là .ts/.tsx và có type rõ ràng.",
        "Page không gọi Axios trực tiếp; gọi custom hook hoặc service của feature.",
        "Không đặt business logic trong router, layout, App.tsx hoặc provider không liên quan.",
        "Component lịch, chart hoặc form nghiệp vụ nằm trong feature sở hữu dữ liệu đó.",
        "DTO backend/frontend lấy từ OpenAPI; không tạo nhiều type response trùng nhau bằng tay.",
        "Mọi màn hình phải kiểm tra desktop và mobile; vùng lịch/biểu đồ có thể cuộn nội bộ nhưng toàn trang không được tràn ngang.",
        "Trước khi đẩy code: chạy npm run typecheck, npm run check:api và npm run build.",
    ])

    add_heading(doc, "11. Các lệnh thường dùng")
    add_code(doc, "npm install                 # cài dependency\nnpm run dev                 # chạy Vite dev server\nnpm run typecheck           # kiểm tra TypeScript\nnpm run build               # typecheck + production build\nnpm run generate:api        # sinh DTO từ Swagger backend\nnpm run generate:api:local  # sinh DTO từ OpenAPI mẫu\nnpm run check:api           # kiểm tra generated DTO có đồng bộ")

    add_heading(doc, "12. Kết luận")
    add_para(doc, "Cấu trúc hiện tại tách rõ vận hành ứng dụng, màn hình, nghiệp vụ và nền tảng dùng chung. TypeScript và OpenAPI giúp kiểm soát contract; React Query và Axios chuẩn hóa dữ liệu; Ant Design chuẩn hóa UI; FullCalendar và Recharts giải quyết các phần hiển thị chuyên biệt. Tuân thủ đúng ownership của app, pages, features và shared sẽ giúp nhóm mở rộng dự án mà không làm code khó tìm hoặc phụ thuộc chéo.")

    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = "Tài liệu cấu trúc Frontend - TuteClass"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        set_run_font(run, 9)

    save_path = DOCX_PATH
    try:
        doc.save(save_path)
    except PermissionError:
        save_path = DOCX_PATH.with_name("Frontend_Structure_TuteClass_updated.docx")
        doc.save(save_path)
    return save_path


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    create_runtime_flow_image()
    create_folder_tree_image()
    create_layer_usage_image()
    create_api_flow_image()
    save_path = build_document()
    print(save_path)


if __name__ == "__main__":
    main()
